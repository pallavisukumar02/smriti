"""Extract text from uploaded files, keeping page numbers for PDFs."""

from __future__ import annotations

import io
import json
from dataclasses import dataclass

SUPPORTED = ("pdf", "docx", "txt", "md", "markdown", "csv", "tsv", "json", "html", "htm", "log", "xml")


@dataclass
class Page:
    text: str
    page: int | None = None


class LoaderError(Exception):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _load_pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [Page((p.extract_text() or ""), i + 1) for i, p in enumerate(reader.pages)]
    if not any(p.text.strip() for p in pages):
        raise LoaderError("No selectable text found. Scanned PDFs need OCR first.")
    return pages


def _load_docx(data: bytes) -> list[Page]:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Page("\n\n".join(parts))]


def _load_html(data: bytes) -> list[Page]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return [Page(soup.get_text("\n"))]


def _load_json(data: bytes) -> list[Page]:
    raw = _decode(data)
    try:
        return [Page(json.dumps(json.loads(raw), indent=2, ensure_ascii=False))]
    except json.JSONDecodeError:
        return [Page(raw)]


def load_file(name: str, data: bytes) -> list[Page]:
    """Return the text of a file as a list of pages."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else "txt"
    if ext not in SUPPORTED:
        raise LoaderError(f"Unsupported file type: .{ext}")
    if ext == "pdf":
        return _load_pdf(data)
    if ext == "docx":
        return _load_docx(data)
    if ext in ("html", "htm"):
        return _load_html(data)
    if ext == "json":
        return _load_json(data)
    return [Page(_decode(data))]
